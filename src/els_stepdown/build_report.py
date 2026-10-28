"""ELS 평가 보고서 생성: results/results.json + 그림 → ELS_몬테카를로_평가_보고서.docx / .pdf
실행 순서: python data.py → python run.py → python build_report.py"""
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core import Product, Market, DAYS

HERE = Path(__file__).resolve().parent
RES = HERE / 'results'
R = json.loads((RES / 'results.json').read_text(encoding='utf-8'))
P, L, B, BL, G, H = R['params'], R['params']['long'], R['base'], R['base_long'], R['greeks'], R['historical']
PROD = Product()
plt.rcParams['font.family'] = 'Malgun Gothic'; plt.rcParams['axes.unicode_minus'] = False
NAVY = '1F3864'
OUT = HERE / 'ELS_몬테카를로_평가_보고서.docx'

# ───────── 추가 그림: 만기 손익 구조, 표본 경로
fig, ax = plt.subplots(figsize=(7.2, 2.9), dpi=170)
x = np.linspace(0.2, 1.3, 400)
c = R['fair_coupon_long']
ax.plot(x * 100, np.where(x >= PROD.barriers[-1], 1 + c * PROD.T, 1 + c * PROD.T) * 100, color='#2F5597', lw=1.6,
        label='낙인 없음: 원금 + 3년 쿠폰')
ax.plot(x * 100, np.where(x >= PROD.barriers[-1], 1 + c * PROD.T, x) * 100, color='#C55A11', lw=1.6, ls='--',
        label='낙인 발생: 만기 배리어 미만이면 성과만큼')
ax.axvline(PROD.barriers[-1] * 100, color='gray', lw=0.8, ls=':'); ax.axvline(PROD.ki * 100, color='gray', lw=0.8, ls=':')
ax.text(PROD.barriers[-1] * 100 + 1, 40, '만기 배리어 75%', fontsize=7.5); ax.text(PROD.ki * 100 + 1, 40, '낙인 50%', fontsize=7.5)
ax.set_xlabel('만기 시점 낮은 쪽 지수 성과 (기준가 대비 %)'); ax.set_ylabel('지급액 (액면 대비 %)')
ax.set_title(f'만기 지급 구조 (조기상환되지 않은 경우, 쿠폰 연 {c * 100:.2f}%)', fontsize=10)
ax.legend(fontsize=7.5, frameon=False, loc='lower right'); ax.grid(alpha=0.25)
[ax.spines[k].set_visible(False) for k in ('top', 'right')]
plt.tight_layout(); plt.savefig(RES / 'fig_payoff.png'); plt.close()

mkt = Market(sigma=np.array([L['sigma_ks200'], L['sigma_spx']]), rho=L['rho_weekly'], r_krw=P['r_krw'], r_usd=P['r_usd'],
             q=np.array(P['q_assumed']), sigma_fx=L['sigma_fx'], rho_spx_fx=L['rho_spx_fx_weekly'])
rng = np.random.default_rng(2026)
steps = int(PROD.T * DAYS); dt = 1 / DAYS
Lc = np.linalg.cholesky(np.array([[1, mkt.rho], [mkt.rho, 1]]))
z = rng.standard_normal((6, steps, 2)) @ Lc.T
perf = np.exp(np.cumsum((mkt.drift() - 0.5 * mkt.sigma ** 2) * dt + mkt.sigma * np.sqrt(dt) * z, axis=1))
worst = perf.min(axis=2)
t = np.arange(1, steps + 1) / DAYS
fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=170)
for i in range(len(worst)):
    ax.plot(t, worst[i] * 100, lw=0.9, alpha=0.85)
ax.step([0] + list(PROD.obs_years), [PROD.barriers[0] * 100] + [b * 100 for b in PROD.barriers], where='post',
        color='black', lw=1.2, label='조기상환 배리어')
ax.axhline(PROD.ki * 100, color='#C00000', lw=1, ls='--', label='낙인 50%')
ax.set_xlabel('경과 연수'); ax.set_ylabel('낮은 쪽 지수 성과 (%)'); ax.set_title('표본 경로 6개 (기본 시나리오, 조기상환 전후 구분 없이 3년 표시)', fontsize=10)
ax.legend(fontsize=7.5, frameon=False); ax.grid(alpha=0.25)
[ax.spines[k].set_visible(False) for k in ('top', 'right')]
plt.tight_layout(); plt.savefig(RES / 'fig_paths.png'); plt.close()

# ───────── 문서 유틸
doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5)
st = doc.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
st.paragraph_format.space_after = Pt(3)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s_ = OxmlElement('w:shd')
    s_.set(qn('w:val'), 'clear'); s_.set(qn('w:color'), 'auto'); s_.set(qn('w:fill'), color); tcPr.append(s_)


def para(text, bold=False, size=None, color=None, after=None, italic=False, align=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    return p


def h(text):
    p = para(text, bold=True, size=12, color=NAVY, after=3); p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '1'), ('color', NAVY)): bt.set(qn(f'w:{k}'), v)
    b.append(bt); pPr.append(b)


def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    p.add_run(text).font.size = Pt(9.5)


def table(rows, widths=None, size=8, first_left=True, hl=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; p = c.paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(size)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_left) else WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, NAVY)
            elif hl and i in hl:
                shade(c, 'DEEAF6')
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pic(name, w=16.5, cap=None):
    doc.add_picture(str(RES / name), width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        para(cap, size=8, color='595959', after=6, align=WD_ALIGN_PARAGRAPH.CENTER)


pc = lambda x, d=1: f'{x * 100:.{d}f}%'
labels = ['6개월', '1년', '1.5년', '2년', '2.5년', '3년(만기 수익)']

# ───────── 표제 + 요약
para('2자산 스텝다운 ELS 몬테카를로 평가 보고서', bold=True, size=17, color=NAVY, after=2)
para(f"KOSPI200 · S&P500 기초 3년 만기 · 평가 기준일 {P['valuation_date']} · 송호빈 (github.com/hopever0105/quant-strategy-lab)",
     size=9, color='595959', after=8)
para('포트폴리오용 개인 분석입니다. 가상의 예시 상품이며 실제 발행 상품의 가격이나 투자 권유가 아닙니다.', size=8, italic=True, color='7F7F7F')

h('요약')
bullet(f"자산평가회사의 주식연계파생상품 평가 절차(발행정보 확인 → 시장데이터 수집 → 파라미터 산출 → 가치평가 → 가격검증)를 그대로 따라 "
       f"2자산 스텝다운 ELS를 Black-Scholes·몬테카를로로 평가했습니다.")
bullet(f"기본 시나리오(2014년 이후 변동성) 공정 쿠폰은 연 {pc(R['fair_coupon_long'], 2)}, 원금손실 확률 {pc(BL['p_loss'])}입니다. "
       f"최근 1년 실현변동성(KOSPI200 {pc(P['sigma_ks200'])})을 넣으면 공정 쿠폰이 연 {pc(R['fair_coupon'], 2)}로 올라, "
       f"같은 상품도 변동성 입력 기간에 따라 가치가 크게 달라졌습니다.")
bullet(f"엔진 검증: 같은 경로 생성기로 계산한 선도가격·콜옵션이 해석해와 표준오차 안에서 일치했고, "
       f"{R['convergence'][-1]['n']:,}개 경로에서 가격 표준오차는 {R['convergence'][-1]['se'] * 100:.3f}%p입니다.")
bullet(f"민감도: 두 지수 변동성 +1%p에 가격 {G['vega_both_1pp'] * 100:+.2f}%, KOSPI200 +1%에 {G['KOSPI200']['delta_1pct'] * 100:+.3f}%. "
       f"상관 변화의 영향은 이 구조에서 작았습니다(+0.05에 {G['corr_0p05'] * 100:+.3f}%).")
bullet(f"과거 실제 경로 {H['n_starts']:,}개(2014-08 ~ 2023-10 발행 가정)에서 6개월 조기상환은 {pc(H['p_k']['0'])}, 원금손실은 "
       f"{pc(H['p_loss'])}로 모형의 위험중립 확률({pc(BL['p_redeem'][0])}, {pc(BL['p_loss'])})과 달랐습니다.")

# ───────── 1. 상품
h('1. 평가 대상 상품 (발행정보)')
table([
    ['항목', '조건'],
    ['기초자산', 'KOSPI200, S&P500 (S&P500 성과도 원화로 지급 = 콴토)'],
    ['만기 / 평가 주기', '3년 / 6개월마다 조기상환 평가 (총 6회)'],
    ['조기상환 배리어', ' - '.join(f'{b * 100:.0f}%' for b in PROD.barriers) + ' (낮은 쪽 지수 성과 기준)'],
    ['낙인(Knock-In)', f'{PROD.ki * 100:.0f}%, 일별 종가 관찰'],
    ['지급', '평가일에 낮은 쪽 성과 ≥ 배리어 → 원금 + 연 쿠폰 × 경과 연수 지급 후 종료'],
    ['만기 미상환 시', '낙인이 한 번도 없었으면 원금 + 3년 쿠폰, 낙인이 있었고 만기 배리어 미만이면 낮은 쪽 성과만큼 지급'],
], [3.6, 13.6], size=8.5)
pic('fig_payoff.png', 14.5)

# ───────── 2. 모형
h('2. 모형과 계산 방법')
para('위험중립 측도에서 두 지수는 상관된 기하 브라운 운동을 따른다고 두었습니다. 원화 투자자 기준이므로 할인은 원화 무위험금리로 하고, '
     '달러 자산인 S&P500의 기대수익률에는 콴토 조정을 넣었습니다.')
table([
    ['구분', '식'],
    ['KOSPI200', 'dS₁/S₁ = (r_KRW − q₁) dt + σ₁ dW₁'],
    ['S&P500 (콴토)', 'dS₂/S₂ = (r_USD − q₂ − ρ(S₂,FX)·σ₂·σ_FX) dt + σ₂ dW₂,   FX = 원/달러'],
    ['상관', 'dW₁·dW₂ = ρ dt  (촐레스키 분해로 상관 난수 생성)'],
    ['가격', 'V = E[ e^(−r_KRW·τ) · (P + c·Y) ],  τ = 지급 시점, P = 원금 부분, Y = 쿠폰이 붙는 연수'],
    ['공정 쿠폰', 'V가 c에 선형이므로 c* = (1 − E[e^(−rτ)P]) / E[e^(−rτ)Y]  → 시뮬레이션 1회로 산출'],
], [3.2, 14.0], size=8.5)
bullet(f"경로: 1년 252일, 3년 {int(PROD.T * DAYS)}일 일별 경로, 기본 {R['convergence'][-2]['n']:,}개. 대조변수(antithetic)로 분산을 줄이고, "
       '표준오차는 대조변수 쌍을 하나의 관측으로 묶어 계산했습니다(쌍끼리 음의 상관이라 그냥 표준편차를 쓰면 과대추정).')
bullet('민감도는 같은 난수(seed 고정)로 입력만 바꿔 재평가했습니다(common random numbers). 델타·감마는 ±1% 중앙차분, 베가는 +1%p.')
pic('fig_paths.png', 14.5)

# ───────── 3. 데이터·파라미터
h('3. 시장데이터와 파라미터')
table([
    ['데이터', '출처', '비고'],
    ['KOSPI200 일별 종가', '네이버 금융 차트(KPI200)', f"2014-08-04 이후 ({P['last']['KS200']:,.2f}, 기준일 종가)"],
    ['S&P500, 원/달러', 'Yahoo Finance (^GSPC, KRW=X)', 'S&P500 1995년~, 원/달러 2003년~'],
    ['원화 무위험금리', P['r_krw_src'], pc(P['r_krw'], 3)],
    ['달러 무위험금리', P['r_usd_src'], pc(P['r_usd'], 2)],
    ['배당수익률', '가정값', f"KOSPI200 {pc(P['q_assumed'][0])}, S&P500 {pc(P['q_assumed'][1])} (민감도로 영향 확인)"],
], [3.8, 6.2, 7.2], size=8.5)
table([
    ['파라미터', '기본: 2014년 이후', '스트레스: 최근 1년'],
    ['산출 기간', L['window'], P['window']],
    ['변동성 KOSPI200', pc(L['sigma_ks200']), pc(P['sigma_ks200'])],
    ['변동성 S&P500', pc(L['sigma_spx']), pc(P['sigma_spx'])],
    ['변동성 원/달러', pc(L['sigma_fx']), pc(P['sigma_fx'])],
    ['상관 KOSPI200–S&P500 (주간)', f"{L['rho_weekly']:.3f}", f"{P['rho_weekly']:.3f}"],
    ['상관 S&P500–원/달러 (주간, 콴토)', f"{L['rho_spx_fx_weekly']:.3f}", f"{P['rho_spx_fx_weekly']:.3f}"],
], [6.2, 5.5, 5.5], size=8.5)
bullet(f"상관계수는 주간 수익률로 구했습니다. 한국 장은 미국 장보다 먼저 닫혀 같은 날짜의 일간 수익률이 서로 다른 정보를 담기 때문에, "
       f"최근 1년 일간 상관은 {P['rho_daily']:.3f}로 주간({P['rho_weekly']:.3f})보다 크게 낮게 나왔습니다(비동시성).")
bullet('변동성은 일간 로그수익률의 표준편차 × √252입니다. 실무에서는 상장 옵션의 내재변동성으로 맞추지만(Calibration), '
       '내재변동성 자료를 구하지 못해 과거 변동성을 두 기간으로 나눠 시나리오로 비교했습니다.')

# ───────── 4. 엔진 검증
h('4. 엔진 검증 (가격검증 1단계)')
para('상품을 평가하기 전에 같은 경로 생성기가 해석해가 있는 상품의 값을 맞히는지 확인했습니다(최근 1년 파라미터, 경로 10만 개).')
rows = [['자산', '선도 해석해', '선도 MC (±SE)', '3년 ATM 콜 해석해', '3년 ATM 콜 MC (±SE)', '차이 / SE']]
for ck in R['check']:
    rows.append([ck['자산'], f"{ck['선도 해석해']:.4f}", f"{ck['선도 MC']:.4f} (±{ck['선도 SE']:.4f})",
                 f"{ck['콜 해석해']:.4f}", f"{ck['콜 MC']:.4f} (±{ck['콜 SE']:.4f})",
                 f"{(ck['콜 MC'] - ck['콜 해석해']) / ck['콜 SE']:+.2f}"])
table(rows, [3.0, 2.2, 3.4, 2.6, 3.6, 2.2], size=8)
bullet('S&P500 콜은 원화 측도의 실효 배당률(r_KRW − 콴토 조정 후 기대수익률)로 Black-Scholes 해석해를 계산해 콴토 drift까지 함께 검증했습니다.')
pic('fig_convergence.png', 14.5, '경로 수에 따른 가격 수렴. 표준오차가 경로 수의 제곱근에 반비례해 줄어듭니다.')
table([['경로 수', '가격 (액면 %)', '표준오차 (%p)']] +
      [[f"{c_['n']:,}", f"{c_['price'] * 100:.3f}", f"{c_['se'] * 100:.3f}"] for c_ in R['convergence']],
      [4.0, 4.0, 4.0], size=8)

# ───────── 5. 결과
h('5. 평가 결과')
table([
    ['항목', '기본: 2014년 이후 변동성', '스트레스: 최근 1년 변동성'],
    ['공정 쿠폰 (연)', pc(R['fair_coupon_long'], 2), pc(R['fair_coupon'], 2)],
    ['가격 표준오차 (공정 쿠폰에서)', f"{BL['se'] * 100:.3f}%p", f"{B['se'] * 100:.3f}%p"],
    ['낙인 확률', pc(BL['p_ki']), pc(B['p_ki'])],
    ['원금손실 확률', pc(BL['p_loss']), pc(B['p_loss'])],
    ['손실 시 평균 손실률', pc(BL['avg_loss_given_loss']), pc(B['avg_loss_given_loss'])],
] + [[f'{lab} 쿠폰 상환 확률', pc(BL['p_redeem'][i]), pc(B['p_redeem'][i])] for i, lab in enumerate(labels)],
    [6.4, 5.4, 5.4], size=8.5, hl=[1])
bullet('공정 쿠폰은 발행사 마진·헤지 비용이 없을 때 가격이 액면과 같아지는 쿠폰입니다. 실제 발행 쿠폰은 여기서 마진만큼 낮아집니다.')
bullet('최근 1년 KOSPI200 변동성이 60%를 넘는 극단 구간이라, 이 값을 3년 내내 유지한다고 보면 낙인 확률이 65%에 이르고 공정 쿠폰이 4배 넘게 올라갑니다. '
       '평가에서 가장 큰 판단은 모형보다 어떤 기간·어떤 종류의 변동성을 넣느냐였습니다.')

# ───────── 6. 민감도
h('6. 민감도 (기본 시나리오)')
table([
    ['입력 변화', '가격 변화 (액면 대비)', '해석'],
    ['KOSPI200 +1%', f"{G['KOSPI200']['delta_1pct'] * 100:+.3f}%", '지수 상승 시 조기상환 가능성↑ → 투자자 가치↑'],
    ['S&P500 +1%', f"{G['S&P500']['delta_1pct'] * 100:+.3f}%", '변동성이 낮아 낮은 쪽이 되는 경우가 적음'],
    ['KOSPI200 감마 (±1%)', f"{G['KOSPI200']['gamma_1pct'] * 100:+.4f}%", '0에 가까움 (몬테카를로 잡음 수준)'],
    ['KOSPI200 변동성 +1%p', f"{G['KOSPI200']['vega_1pp'] * 100:+.3f}%", '투자자는 변동성 매도 포지션'],
    ['S&P500 변동성 +1%p', f"{G['S&P500']['vega_1pp'] * 100:+.3f}%", ''],
    ['두 지수 변동성 +1%p', f"{G['vega_both_1pp'] * 100:+.3f}%", '가장 큰 위험 요인'],
    ['상관 +0.05', f"{G['corr_0p05'] * 100:+.3f}%", '이 구조에서는 작음'],
    ['원화 금리 +1bp', f"{G['rate_1bp'] * 100:+.4f}%", ''],
], [4.6, 3.8, 8.8], size=8.5)
pic('fig_sens.png', 16, '입력값에 따른 공정 쿠폰 (경로 5만 개, 공정 쿠폰의 몬테카를로 오차 약 ±0.1%p)')
S = R['sensitivity']
table([['KOSPI200 배당수익률', *[pc(q, 1) for q, _ in S['q_ks200']]],
       ['공정 쿠폰 (연)', *[pc(v, 2) for _, v in S['q_ks200']]]], size=8)
bullet('상관을 0에서 0.8까지 바꿔도 공정 쿠폰 변화가 0.13%p 안이라 몬테카를로 오차와 비슷한 크기였습니다. '
       'KOSPI200 변동성이 S&P500보다 커서 대부분의 경로에서 KOSPI200이 낮은 쪽이 되기 때문으로 해석했습니다.')
bullet('배당수익률은 가정값이라 따로 보였습니다. 0~3% 범위에서 공정 쿠폰이 약 1.4%p 움직여, 실제 평가에서는 배당 선물·추정치를 입력으로 써야 합니다.')

# ───────── 7. 과거 경로
h('7. 과거 실제 경로와의 비교 (가격검증 2단계)')
para(f"{H['first']}부터 {H['last']}까지 KOSPI200 거래일마다 같은 상품을 발행했다고 보고, 실제 종가로 상환 시점을 계산했습니다"
     f"({H['n_starts']:,}개, 서로 겹치는 3년 구간, 환율 효과 제외, 쿠폰은 기본 시나리오 공정 쿠폰).")
pic('fig_redeem.png', 15.5)
table([['', *labels, '원금손실'],
       ['몬테카를로 (기본)', *[pc(v) for v in BL['p_redeem']], pc(BL['p_loss'])],
       ['몬테카를로 (스트레스)', *[pc(v) for v in B['p_redeem']], pc(B['p_loss'])],
       ['과거 실제 경로', *[pc(H['p_k'].get(str(i), 0)) for i in range(6)], pc(H['p_loss'])]],
      [3.4, 1.8, 1.7, 1.7, 1.7, 1.7, 2.4, 1.8], size=7.5)
bullet(f"과거 경로는 6개월 상환이 더 많고 손실이 없었습니다. 위험중립 확률은 가격을 매기기 위한 확률이라 실제 빈도와 같을 이유가 없고, "
       f"과거 표본은 서로 겹치는 3년 구간이라 독립 관측 수는 실제로 3~4개 수준입니다. 결과 차이를 모형 오류로 보지 않고 두 확률의 성격 차이로 기록했습니다.")
bullet(f"과거 경로의 평균 연환산 수익률은 {pc(H['mean_ann_ret'], 2)}로 공정 쿠폰과 비슷했습니다(대부분 6개월 조기상환).")

# ───────── 8. 한계
h('8. 한계와 다음 단계')
bullet('변동성: 내재변동성·변동성 스마일 대신 상수 과거 변동성을 썼습니다. 다음 단계는 KOSPI200·S&P500 옵션 내재변동성으로 Calibration 하는 것입니다.')
bullet('배당·금리: 배당수익률은 가정값, 금리는 3년 만기 단일 금리(평탄 커브)입니다.')
bullet('모형: 상수 변동성 GBM이라 급락 구간의 변동성 상승(레버리지 효과)을 반영하지 못해 낙인 확률을 낮게 볼 수 있습니다. 국지변동성·확률변동성 모형 비교가 다음 과제입니다.')
bullet('데이터: KOSPI200 일별 데이터가 2014-08부터라 과거 경로 비교에 2008년 같은 급락 구간이 없습니다. 휴장일이 달라 S&P500은 직전 종가로 채웠습니다.')

# ───────── 9. 재현
h('9. 재현 방법')
table([
    ['파일', '내용'],
    ['data.py', '시장데이터 수집 (네이버·Yahoo·ECOS·FRED, KRX 요청 없음) → data/'],
    ['core.py', '몬테카를로 엔진(상관 GBM, 대조변수, 콴토), 공정 쿠폰, 해석해 검증'],
    ['run.py', '파라미터 산출 → 엔진 검증 → 수렴 → 민감도 → 과거 경로 → results/results.json, 그림'],
    ['build_report.py', '이 보고서 생성'],
], [3.6, 13.6], size=8.5)
para('python data.py → python run.py (약 4분) → python build_report.py', size=9)

doc.save(OUT)

# PDF (Word가 있으면 Word로 변환)
try:
    import win32com.client
    w = win32com.client.DispatchEx('Word.Application'); w.Visible = False; w.DisplayAlerts = 0
    try:
        d = w.Documents.Open(str(OUT), ReadOnly=True)
        d.ExportAsFixedFormat(str(OUT.with_suffix('.pdf')), 17)
        print('쪽수:', d.ComputeStatistics(2))
        d.Close(False)
    finally:
        w.Quit()
except Exception as e:
    print('PDF 변환 실패:', e)
print('완료:', OUT)
