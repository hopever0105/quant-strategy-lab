"""전환사채 LSMC 평가 보고서 생성: results/results.json + 그림 → 전환사채_LSMC_평가_보고서.docx / .pdf
실행 순서: python data.py → python run.py → python build_report.py"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
RES = HERE / 'results'
R = json.loads((RES / 'results.json').read_text(encoding='utf-8'))
P, C, D, G, S, PR = R['params'], R['check'], R['decompose'], R['greeks'], R['scenarios'], R['product']
E, PT, NR = C['european'], C['put'], C['no_refix']
FULL = D['full_detail']
NAVY = '1F3864'
OUT = HERE / '전환사채_LSMC_평가_보고서.docx'
pct = lambda x, d=1: f'{x * 100:.{d}f}%'
f2 = lambda x: f'{x:.2f}'

doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5)
st = doc.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')
st.paragraph_format.space_after = Pt(3)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); s_ = OxmlElement('w:shd')
    s_.set(qn('w:val'), 'clear'); s_.set(qn('w:color'), 'auto'); s_.set(qn('w:fill'), color); tcPr.append(s_)


def para(text, bold=False, size=None, color=None, after=None, align=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    return p


def h(text):
    p = para(text, bold=True, size=12, color=NAVY, after=3); p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '1'), ('color', NAVY)): bt.set(qn(f'w:{k}'), v)
    b.append(bt); pPr.append(b)


def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    p.add_run(text).font.size = Pt(9.5)


def table(rows, widths=None, size=8, hl=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; p = c.paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(size)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); shade(c, NAVY)
            elif hl and i in hl:
                shade(c, 'DEEAF6')
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pic(name, w=16.5, cap=None):
    doc.add_picture(str(RES / name), width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        para(cap, size=8, color='595959', after=6, align=WD_ALIGN_PARAGRAPH.CENTER)


# ───────── 1쪽: 요약
para('전환사채(CB) LSMC 평가', bold=True, size=18, color=NAVY, after=2)
para(f"리픽싱·조기상환청구권이 있는 3년 만기 CB (가상의 예시 상품) | 평가 기준일 {P['valuation_date']} | 포트폴리오용 개인 분석, 투자 권유 아님",
     size=9, color='595959', after=8)
h('요약')
bullet(f"국내 CB의 미국형 전환권, 조기상환청구권(풋), 리픽싱을 모두 넣은 가상 CB를 LSMC로 평가했다. CB 가치는 액면 100 기준 {f2(D['full'])}"
       f"(±{FULL['se']:.2f}), 전환 {pct(FULL['p_conv'])} · 풋 {pct(FULL['p_put'])} · 만기상환 {pct(FULL['p_mat'])}.")
bullet(f"가치 분해: 순수 채권 {f2(D['bond_no_put'])} + 전환권 {f2(D['v_conversion'])} + 조기상환청구권 {f2(D['v_put'])} + 리픽싱 {f2(D['v_refix'])}.")
bullet(f"엔진 검증: 만기 전환 CB 해석해 {f2(E['closed'])} vs MC {f2(E['mc'])}(±{E['mc_se']:.2f}), 버뮤단 풋 트리 {PT['tree']:.4f} vs LSMC {PT['lsmc']:.4f}, "
       f"리픽싱 없는 CB 트리 {f2(NR['tree'])} vs LSMC {f2(NR['euro']['oos'])}.")
bullet(f"회귀 기저가 결과를 바꿨다: 패리티 다항식만 쓰면 트리보다 {NR['tree'] - NR['cubic']['oos']:.2f} 낮았고(LSMC의 아래쪽 편향), "
       f"유럽형 CB 해석해를 기저에 넣자 차이가 {abs(NR['tree'] - NR['euro']['oos']):.2f}로 줄었다.")

h('상품 조건 (가상)')
puts = ', '.join(f"{k} {v:.2f}" for k, v in PR['puts'].items())
table([
    ['항목', '조건'],
    ['기초주식 / 최초 전환가액', f"{P['name']} ({P['ticker']}), 기준일 종가 {P['s0_krw']:,.0f}원 = 최초 전환가액 (할증 0%)"],
    ['만기 / 표면이자 / 만기보장수익률', f"3년 / 0% / 2% (3개월 복리), 만기상환 {PR['redemption']:.2f}"],
    ['전환', '1년 후 ~ 만기 1개월 전, 매주(모형상 주 단위 행사)'],
    ['조기상환청구권(풋)', f"1.5년 후 3개월마다, 조기상환수익률 2%: {puts}"],
    ['리픽싱', '3개월마다, 주가 < 전환가액이면 주가로 하향. 하한 최초 전환가액의 70%'],
    ['넣지 않은 것', '발행사 매도청구권(콜), 희석, 리픽싱 기준가 가중평균 산식'],
], [4.6, 12.8], size=8)

h('모형과 파라미터')
table([
    ['구분', '값', '출처 / 근거'],
    ['주가 과정', '위험중립 GBM, 주 단위 정확 이산화', 'drift = r − q'],
    ['변동성 (기본)', pct(P['sigma_3y']), f"최근 3년 일간 로그수익률 ({P['n_daily'][0]}일), 만기와 같은 기간"],
    ['변동성 (스트레스)', pct(P['sigma_1y']), f"최근 1년 ({P['n_daily'][1]}일). 참고: 전체 {pct(P['sigma_all'])} ({P['window_all']})"],
    ['무위험금리', pct(P['r'], 3), P['r_src']],
    ['신용스프레드 (기본 / 스트레스)', f"{pct(P['s_aa'], 2)}p / {pct(P['s_bbb'], 2)}p", P['spread_src']],
    ['배당수익률', pct(P['q'], 2), P['q_src']],
    ['신용위험 처리', 'Tsiveriotis-Fernandes', '전환분은 r, 현금 상환분(만기·풋)은 r + s로 할인'],
    ['LSMC', '훈련 10만 + 평가 20만 경로', '회귀계수와 가격을 다른 난수로(out-of-sample), 대조변수'],
], [4.2, 4.6, 8.6], size=8)

# ───────── 2쪽: 검증
doc.add_page_break()
h('엔진 검증')
table([
    ['검증', '기준값', 'LSMC / MC', '확인한 것'],
    ['① 만기에만 전환하는 CB', f"해석해 {f2(E['closed'])} / 트리 {f2(E['tree'])}", f"{f2(E['mc'])} (±{E['mc_se']:.2f})", '경로 생성·TF 할인'],
    ['② 주식 버뮤단 풋(1년, ATM)', f"트리 {PT['tree']:.4f} (유럽형 {PT['european']:.4f})", f"{PT['lsmc']:.4f} (±{PT['lsmc_se']:.4f})", '회귀·조기행사 로직'],
    ['③ 리픽싱 없는 CB: 다항식 기저', f"트리 {f2(NR['tree'])}", f"{f2(NR['cubic']['oos'])} (±{NR['cubic']['se']:.2f})", '기저가 부족하면 아래로 편향'],
    ['③ 리픽싱 없는 CB: 해석해 기저', f"트리 {f2(NR['tree'])}", f"{f2(NR['euro']['oos'])} (±{NR['euro']['se']:.2f})", '본 평가에 쓴 기저'],
], [4.6, 4.8, 3.6, 4.4], size=8, hl=[4])
para('이항트리는 한 주를 여러 스텝으로 나누되 행사는 주 단위 시점에서만 허용해 LSMC와 행사 기회를 맞췄다. '
     'TF 트리는 노드마다 전체 가치와 현금 부분을 함께 들고 가며 서로 다른 할인율을 적용한다. '
     '트리 스텝 수에 따른 값: ' + ', '.join(f"한 주 {m}스텝 {v:.2f}" for m, v in C['tree_steps']) + '.', size=9)
para(f"LSMC는 회귀로 추정한 계속 보유 가치로 행사 여부를 정하므로 회귀가 틀리면 덜 최적인 행사를 하게 되고, 그 결과는 참값보다 낮아진다. "
     f"in-sample(회귀와 가격을 같은 경로로) 값은 {f2(NR['euro']['in_sample'])}, out-of-sample 값은 {f2(NR['euro']['oos'])}였다.", size=9)
pic('fig_convergence.png', 15, '평가 경로 수에 따른 CB 가치 수렴')

# ───────── 3쪽: 결과
doc.add_page_break()
h('가치 분해 (액면 100 기준)')
table([
    ['구성', '가치', '설명'],
    ['순수 채권 (만기상환만)', f2(D['bond_no_put']), '만기상환액을 r + s로 할인'],
    ['+ 전환권', f"+{f2(D['v_conversion'])}", f"전환만 있는 CB {f2(D['conv_only'])}"],
    ['+ 조기상환청구권', f"+{f2(D['v_put'])}", f"전환 + 풋 {f2(D['conv_put'])}"],
    ['+ 리픽싱', f"+{f2(D['v_refix'])}", '전환가액 하향 조정'],
    ['CB 가치', f2(D['full']), f"채권 가치(전환권 없음, 풋 포함) {f2(D['bond_with_put'])}"],
], [5.0, 3.0, 9.4], size=8, hl=[5])
pic('fig_parity.png', 15.5, '주가에 따른 CB 가치 (같은 난수, 경로 10만 개)')
pic('fig_paths.png', 15.5, '표본 경로와 리픽싱된 전환가액')

# ───────── 4쪽: 시나리오·민감도
doc.add_page_break()
h('시나리오')
rows = [['시나리오', 'CB 가치', '채권 가치(풋 포함)', '전환', '풋', '만기상환']]
for k, v in S.items():
    rows.append([k, f2(v['price']), f2(v['bond_with_put']), pct(v['p_conv']), pct(v['p_put']), pct(v['p_mat'])])
table(rows, [5.2, 2.2, 3.2, 2.2, 2.2, 2.4], size=8)
h('민감도 (같은 난수로 재평가)')
table([
    ['입력 변화', '가치 변화 (액면 100 기준)'],
    ['주가 +1% (델타)', f"{G['delta_1pct']:+.3f}"],
    ['주가 ±1% 감마', f"{G['gamma_1pct']:+.3f}"],
    ['변동성 +1%p', f"{G['vega_1pp']:+.3f}"],
    ['신용스프레드 +50bp', f"{G['spread_50bp']:+.3f}"],
    ['무위험금리 +10bp', f"{G['rate_10bp']:+.3f}"],
    ['배당수익률 +1%p', f"{G['q_1pp']:+.3f}"],
], [8.7, 8.7], size=8)
pic('fig_sens.png', 16.5, '변동성·신용스프레드·리픽싱 하한에 따른 CB 가치')

h('한계와 다음 단계')
bullet('상수 과거 변동성을 썼다. 실무라면 상장 옵션이 없는 개별 주식의 변동성을 어떻게 정할지(과거 변동성 기간, 유사 종목)가 가장 큰 가정이다.')
bullet('신용스프레드는 발행사 개별 등급이 아니라 회사채 AA-·BBB- 평균 스프레드로 시나리오를 나눴다.')
bullet('리픽싱 기준가는 실제 약정의 가중평균 산식 대신 리픽싱일 주가로 단순화했고, 발행사 매도청구권과 희석은 넣지 않았다.')
bullet('LSMC 결과는 회귀 기저에 따라 아래로 치우칠 수 있다. 리픽싱이 있는 경우는 트리 기준값이 없어 같은 기저의 편향이 리픽싱 없는 경우와 비슷하다고 가정했다.')

doc.save(OUT)
print('저장', OUT)
try:
    import win32com.client
    w = win32com.client.DispatchEx('Word.Application'); w.Visible = False; w.DisplayAlerts = 0
    try:
        d = w.Documents.Open(str(OUT), ReadOnly=True)
        d.ExportAsFixedFormat(str(OUT.with_suffix('.pdf')), 17)
        print('쪽수', d.ComputeStatistics(2)); d.Close(False)
    finally:
        w.Quit()
except Exception as ex:
    print('PDF 변환 생략:', ex)
